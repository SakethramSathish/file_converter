"""
Document conversion service.
Handles: DOCX↔PDF, DOCX↔TXT, ODT→DOCX, ODT→PDF
"""
import os
import tempfile
from docx import Document
from docx2pdf import convert as docx2pdf_convert
from pdf2docx import Converter


def docx_to_pdf(input_path: str, output_dir: str) -> str:
    """Convert DOCX to PDF using docx2pdf."""
    output_path = os.path.join(output_dir, "output.pdf")
    docx2pdf_convert(input_path, output_path)
    return output_path


def pdf_to_docx(input_path: str, output_dir: str) -> str:
    """Convert PDF to DOCX using pdf2docx."""
    output_path = os.path.join(output_dir, "output.docx")
    cv = Converter(input_path)
    cv.convert(output_path, start=0, end=None)
    cv.close()
    return output_path


def txt_to_docx(input_path: str, output_dir: str) -> str:
    """Convert TXT to DOCX by creating paragraphs from each line."""
    output_path = os.path.join(output_dir, "output.docx")
    doc = Document()
    with open(input_path, "r", encoding="utf-8") as f:
        for line in f:
            doc.add_paragraph(line.strip())
    doc.save(output_path)
    return output_path


def docx_to_txt(input_path: str, output_dir: str) -> str:
    """Convert DOCX to TXT by extracting paragraph text."""
    output_path = os.path.join(output_dir, "output.txt")
    doc = Document(input_path)
    with open(output_path, "w", encoding="utf-8") as f:
        for para in doc.paragraphs:
            f.write(para.text + "\n")
    return output_path


def odt_to_docx(input_path: str, output_dir: str) -> str:
    """Convert ODT to DOCX using LibreOffice headless."""
    os.system(f'libreoffice --headless --convert-to docx "{input_path}" --outdir "{output_dir}"')
    base = os.path.splitext(os.path.basename(input_path))[0]
    output_path = os.path.join(output_dir, f"{base}.docx")
    if not os.path.exists(output_path):
        raise FileNotFoundError("LibreOffice conversion failed. Ensure LibreOffice is installed.")
    return output_path


def odt_to_pdf(input_path: str, output_dir: str) -> str:
    """Convert ODT to PDF using LibreOffice headless."""
    os.system(f'libreoffice --headless --convert-to pdf "{input_path}" --outdir "{output_dir}"')
    base = os.path.splitext(os.path.basename(input_path))[0]
    output_path = os.path.join(output_dir, f"{base}.pdf")
    if not os.path.exists(output_path):
        raise FileNotFoundError("LibreOffice conversion failed. Ensure LibreOffice is installed.")
    return output_path
